import csv
import io
import os
import re
import unicodedata
from datetime import date, datetime
from typing import Any
from sqlalchemy import and_, case, distinct, false, func, or_, true

import pandas as pd
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import Column, Date, Integer, MetaData, Numeric, String, Table, Text, cast, delete, insert, inspect, select, text, update

from ..core.database import engine
from ..models.user import User

LOAN_TABLE_NAME = 'loan_imports'
ACTIVITY_TABLE_NAME = 'loan_activities_status'
BATCH_SIZE = 500
REQUIRED_IMPORT_COLUMNS = {
    'BRCD', 'CUSTSEQ', 'CUSTNM', 'DISBURSEMENT_AMOUNT', 'DU_NO',
    'LOAN_TYPE', 'CUSTOMER_TYPE_CODE', 'ADDR1', 'DSBSDT', 'INTTRMMTH',
}
MATURITY_DATE_HEADERS = ('APPRMATDT', 'DSBSMATDT', 'DUE_DATE', 'MATURITY_DATE', 'NGAY_DEN_HAN', 'NGAY_TO_HAN', 'NGAY_DAO_HAN')
MAPPED_HEADER_KEYS = {
    'BRCD', 'CUSTSEQ', 'CUSTNM', 'DISBURSEMENT_AMOUNT', 'DU_NO',
    'LOAN_TYPE', 'CUSTOMER_TYPE_CODE', 'ADDR1', 'DSBSDT', 'DSBSMATDT', 'APPRMATDT', 'INTTRMMTH',
    'NHOM_NO', 'LAST_REPAY_DATE', 'NEXT_REPAY_DATE', 'PASTDUE_INTEREST_AMOUNT',
    'TOTAL_INTEREST_REPAY_AMOUNT',
}
HEADER_LABELS = {
    'BRCD': 'Mã chi nhánh', 'CUSTSEQ': 'Mã khách hàng', 'CUSTNM': 'Tên khách hàng',
    'DISBURSEMENT_AMOUNT': 'Số tiền vay', 'DU_NO': 'Dư nợ gốc', 'LOAN_TYPE': 'Loại khoản vay',
    'ADDR1': 'Địa chỉ', 'DSBSDT': 'Ngày bắt đầu vay', 'DSBSMATDT': 'Ngày đáo hạn giải ngân', 'APPRMATDT': 'Ngày đáo hạn phê duyệt',
    'DSBSSEQ': 'Mã giải ngân',
    'INTTRMMTH': 'Thời hạn', 'CUSTOMER_TYPE_CODE': 'Mã loại khách hàng', 'NHOM_NO': 'Nhóm nợ', 'LAST_REPAY_DATE': 'Ngày trả nợ gần nhất',
    'NEXT_REPAY_DATE': 'Ngày trả nợ tiếp theo',
    'PASTDUE_INTEREST_AMOUNT': 'Lãi quá hạn',
    'TOTAL_INTEREST_REPAY_AMOUNT': 'Tổng tiền lãi phải trả',
}

metadata = MetaData()
activity_metadata = MetaData()
loan_table: Table | None = None
activity_table: Table | None = None
original_headers: list[str] = []
column_keys: list[str] = []
loan_due_date_column: str | None = None
loan_dsbsseq_column: str | None = None


def normalize_header(value: Any) -> str:
    return str(value).strip()


def _normalize_search_text(value: str) -> str:
    if value is None:
        return ''
    normalized = unicodedata.normalize('NFD', str(value).strip().lower())
    return ''.join(ch for ch in normalized if unicodedata.category(ch) != 'Mn')


def _ensure_unaccent_extension() -> None:
    try:
        with engine.begin() as connection:
            connection.execute(text("CREATE EXTENSION IF NOT EXISTS unaccent"))
    except Exception:
        pass


def _normalized_like_expression(column: Any, search_value: str):
    normalized_value = _normalize_search_text(search_value)
    if not normalized_value:
        return None
    _ensure_unaccent_extension()
    return func.lower(func.unaccent(cast(column, String))).like(f"%{normalized_value}%")


def sanitize_column_name(value: Any) -> str:
    cleaned = re.sub(r'[^0-9a-zA-Z_]', '_', normalize_header(value))
    cleaned = re.sub(r'_+', '_', cleaned).strip('_') or 'column'
    return f'col_{cleaned}' if cleaned[0].isdigit() else cleaned


def _configure_headers(headers: list[str]) -> None:
    global original_headers, column_keys, loan_due_date_column, loan_dsbsseq_column
    original_headers = [normalize_header(header) for header in headers]
    used: set[str] = set()
    column_keys = []
    for header in original_headers:
        base = sanitize_column_name(header)
        key = base
        suffix = 2
        while key in used:
            key = f'{base}_{suffix}'
            suffix += 1
        used.add(key)
        column_keys.append(key)
    loan_dsbsseq_column = next((key for header, key in zip(original_headers, column_keys) if header.upper() == 'DSBSSEQ'), None)
    loan_due_date_column = None
    candidates = (*MATURITY_DATE_HEADERS, 'DSBSDT')
    for candidate in candidates:
        for header, key in zip(original_headers, column_keys):
            if header.upper() == candidate or candidate.lower() in header.lower():
                loan_due_date_column = key
                break
        if loan_due_date_column:
            break


def _ensure_activity_table() -> Table:
    global activity_table
    inspector = inspect(engine)
    if inspector.has_table(ACTIVITY_TABLE_NAME):
        existing = {column['name'] for column in inspector.get_columns(ACTIVITY_TABLE_NAME)}
        if 'gdv' not in existing:
            with engine.begin() as connection:
                connection.execute(text(f'ALTER TABLE {ACTIVITY_TABLE_NAME} ADD COLUMN gdv VARCHAR(50)'))
        activity_table = Table(ACTIVITY_TABLE_NAME, activity_metadata, autoload_with=engine, extend_existing=True)
        return activity_table

    activity_table = Table(
        ACTIVITY_TABLE_NAME, activity_metadata,
        Column('id', Integer, primary_key=True),
        Column('dsbsseq', String(100), nullable=False, unique=True),
        Column('status', String(50)), Column('gdv_note', Text), Column('gdv', String(50)),
        Column('updated_at', String(50)),
    )
    activity_metadata.create_all(engine)
    return activity_table


def _ensure_loan_table(headers: list[str] | None = None) -> Table | None:
    global loan_table
    inspector = inspect(engine)
    headers_were_supplied = headers is not None
    if headers is None:
        if not inspector.has_table(LOAN_TABLE_NAME):
            return None
        headers = [column['name'] for column in inspector.get_columns(LOAN_TABLE_NAME)
                   if column['name'].lower() not in {'id', 'status', 'gdv_note', 'source_file', 'imported_at'}]
    if not headers:
        return None

    # For an import, follow main1.py: the uploaded headers define the table
    # schema. Reuse the table only when its dynamic columns match exactly.
    if headers_were_supplied and inspector.has_table(LOAN_TABLE_NAME):
        _configure_headers(headers)
        expected = {'id', 'status', 'gdv_note', 'source_file', 'imported_at', *column_keys}
        existing = {column['name'] for column in inspector.get_columns(LOAN_TABLE_NAME)}
        if existing != expected:
            with engine.begin() as connection:
                connection.execute(text(f'DROP TABLE IF EXISTS {LOAN_TABLE_NAME}'))
            if LOAN_TABLE_NAME in metadata.tables:
                metadata.remove(metadata.tables[LOAN_TABLE_NAME])
            loan_table = None
        else:
            loan_table = Table(LOAN_TABLE_NAME, metadata, autoload_with=engine, extend_existing=True)
            _ensure_activity_table()
            return loan_table

    # When listing existing data, extend legacy loan_imports in place with
    # technical columns instead of recreating it.
    if not headers_were_supplied and inspector.has_table(LOAN_TABLE_NAME):
        existing = {column['name'].lower() for column in inspector.get_columns(LOAN_TABLE_NAME)}
        additions = []
        if 'id' not in existing:
            additions.append(f'ALTER TABLE {LOAN_TABLE_NAME} ADD COLUMN id SERIAL')
        if 'status' not in existing:
            additions.append(f'ALTER TABLE {LOAN_TABLE_NAME} ADD COLUMN status VARCHAR(50)')
        if 'gdv_note' not in existing:
            additions.append(f'ALTER TABLE {LOAN_TABLE_NAME} ADD COLUMN gdv_note TEXT')
        if 'source_file' not in existing:
            additions.append(f'ALTER TABLE {LOAN_TABLE_NAME} ADD COLUMN source_file VARCHAR(255)')
        if 'imported_at' not in existing:
            additions.append(f'ALTER TABLE {LOAN_TABLE_NAME} ADD COLUMN imported_at VARCHAR(50)')
        if additions:
            with engine.begin() as connection:
                for statement in additions:
                    connection.execute(text(statement))
        inspector = inspect(engine)
        physical_columns = [column['name'] for column in inspector.get_columns(LOAN_TABLE_NAME)]
        headers = [column for column in physical_columns if column.lower() not in {'id', 'status', 'gdv_note', 'source_file', 'imported_at'}]

    _configure_headers(headers)
    if inspect(engine).has_table(LOAN_TABLE_NAME):
        loan_table = Table(LOAN_TABLE_NAME, metadata, autoload_with=engine, extend_existing=True)
        _ensure_activity_table()
        return loan_table
    columns = [
        Column('id', Integer, primary_key=True),
        Column('status', String(50), nullable=False), Column('gdv_note', Text),
        Column('source_file', String(255)), Column('imported_at', String(50)),
        *(Column(key, Text) for key in column_keys),
    ]
    loan_table = Table(LOAN_TABLE_NAME, metadata, *columns, extend_existing=True)
    metadata.create_all(engine)
    loan_table = Table(LOAN_TABLE_NAME, metadata, autoload_with=engine)
    _ensure_activity_table()
    return loan_table


def determine_status(raw_value: Any) -> str:
    if raw_value is None or (isinstance(raw_value, float) and pd.isna(raw_value)):
        return 'Đang vay'
    normalized = str(raw_value).strip().replace(' ', '')
    if not normalized:
        return 'Đang vay'
    if ',' in normalized and '.' in normalized:
        normalized = normalized.replace('.', '').replace(',', '.') if normalized.rfind(',') > normalized.rfind('.') else normalized.replace(',', '')
    else:
        normalized = normalized.replace(',', '').replace('.', '')
    cleaned = re.sub(r'[^0-9.-]', '', normalized)
    try:
        return 'Đã tất toán' if float(cleaned or 1) == 0 else 'Đang vay'
    except ValueError:
        return 'Đang vay'


def _parse_date(value: Any) -> date | None:
    if value is None or str(value).strip() == '':
        return None
    try:
        parsed = pd.to_datetime(value)
        return parsed.date()
    except (TypeError, ValueError):
        return None


def _format_value(key: str, value: Any) -> Any:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    upper = key.upper()
    if upper in {'DSBSDT', 'DSBSMATDT', 'APPRMATDT', 'LAST_REPAY_DATE', 'NEXT_REPAY_DATE'}:
        parsed = _parse_date(value)
        return parsed.strftime('%d/%m/%Y') if parsed else ''
    if upper in {'DISBURSEMENT_AMOUNT', 'DU_NO'}:
        text_value = str(value).strip().replace(' ', '')
        text_value = text_value.replace(',', '') if ',' in text_value and '.' not in text_value else text_value.replace('.', '').replace(',', '.') if ',' in text_value else text_value.replace('.', '')
        try:
            return f'{float(re.sub(r"[^0-9.-]", "", text_value)):,.0f}'
        except ValueError:
            return str(value)
    return value


def _activity_by_key(connection, key: str | None) -> dict[str, Any] | None:
    if not key or activity_table is None or not key.strip():
        return None
    return connection.execute(
        select(
            activity_table.c.dsbsseq,
            activity_table.c.status,
            activity_table.c.gdv_note,
            activity_table.c.gdv,
            User.fullname.label('gdv_name'),
        )
        .select_from(activity_table.outerjoin(User.__table__, activity_table.c.gdv == User.__table__.c.username))
        .where(activity_table.c.dsbsseq == key.strip())
    ).mappings().first()


def _activities_by_keys(connection, keys: set[str]) -> dict[str, dict[str, Any]]:
    if not keys or activity_table is None:
        return {}
    rows = connection.execute(
        select(
            activity_table.c.dsbsseq,
            activity_table.c.status,
            activity_table.c.gdv_note,
            activity_table.c.gdv,
            User.fullname.label('gdv_name'),
        )
        .select_from(activity_table.outerjoin(User.__table__, activity_table.c.gdv == User.__table__.c.username))
        .where(activity_table.c.dsbsseq.in_(keys))
    ).mappings().all()
    return {str(row['dsbsseq']).strip(): dict(row) for row in rows}


def _resolved_status_expression(table: Table):
    if activity_table is None or loan_dsbsseq_column is None:
        return table.c.status
    return func.coalesce(activity_table.c.status, table.c.status)


def _outstanding_amount_expression(table: Table):
    amount_key = _key_for_header('DU_NO')
    if not amount_key or amount_key not in table.c:
        return cast(0, Numeric)
    raw_amount = func.regexp_replace(cast(table.c[amount_key], String), '[^0-9,.-]', '', 'g')
    normalized_amount = case(
        (and_(func.strpos(raw_amount, ',') > 0, func.strpos(raw_amount, '.') > 0,
              func.strpos(func.reverse(raw_amount), ',') < func.strpos(func.reverse(raw_amount), '.')),
         func.replace(func.replace(raw_amount, '.', ''), ',', '.')),
        (and_(func.strpos(raw_amount, ',') > 0, func.strpos(raw_amount, '.') > 0), func.replace(raw_amount, ',', '')),
        (func.strpos(raw_amount, ',') > 0, func.replace(raw_amount, ',', '')),
        else_=func.replace(raw_amount, '.', ''),
    )
    return cast(func.nullif(normalized_amount, ''), Numeric)


def _customer_type_expression(table: Table):
    customer_type_key = _key_for_header('CUSTOMER_TYPE_CODE')
    if not customer_type_key or customer_type_key not in table.c:
        return None
    customer_type_code = func.trim(cast(table.c[customer_type_key], String))
    return case(
        (customer_type_code == '100', 'individual'),
        (customer_type_code.in_(['520', '530']), 'legal_entity'),
        else_='unknown',
    )


def _maturity_date_expression(table: Table):
    apprmatdt_key = _key_for_header('APPRMATDT')
    if apprmatdt_key and apprmatdt_key in table.c:
        return cast(table.c[apprmatdt_key], Date)
    if loan_due_date_column and loan_due_date_column in table.c:
        return cast(table.c[loan_due_date_column], Date)
    return None


def _key_for_header(header: str) -> str | None:
    normalized = header.strip().upper()
    return next((key for original, key in zip(original_headers, column_keys) if original.upper() == normalized), None)


def _report_number(value: Any) -> float:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return 0
    normalized = str(value).strip().replace(' ', '')
    if not normalized:
        return 0
    if ',' in normalized and '.' in normalized:
        normalized = normalized.replace('.', '').replace(',', '.') if normalized.rfind(',') > normalized.rfind('.') else normalized.replace(',', '')
    elif ',' in normalized:
        normalized = normalized.replace(',', '')
    try:
        return float(re.sub(r'[^0-9.-]', '', normalized) or 0)
    except ValueError:
        return 0


def _report_date(value: Any) -> date | None:
    if value is None or str(value).strip() == '':
        return None
    parsed = pd.to_datetime(value, errors='coerce')
    return None if pd.isna(parsed) else parsed.date()


def _is_closed_loan(row: dict[str, Any]) -> bool:
    amount_keys = (
        _key_for_header('DU_NO'),
        _key_for_header('PASTDUE_INTEREST_AMOUNT'),
        _key_for_header('TOTAL_INTEREST_REPAY_AMOUNT'),
    )
    return all(key is not None for key in amount_keys) and all(_report_number(row.get(key)) == 0 for key in amount_keys)


def _is_active_loan_at(
    row: dict[str, Any],
    cutoff: date,
    disbursement_key: str | None,
    maturity_key: str | None,
) -> bool:
    disbursement_date = _report_date(row.get(disbursement_key)) if disbursement_key else None
    if disbursement_date is None or disbursement_date > cutoff:
        return False
    maturity_date = _report_date(row.get(maturity_key)) if maturity_key else None
    return not (_is_closed_loan(row) and maturity_date is not None and maturity_date <= cutoff)


def _missing_import_columns(headers: list[str]) -> list[str]:
    normalized_headers = {header.upper() for header in headers}
    missing = REQUIRED_IMPORT_COLUMNS - normalized_headers
    if not normalized_headers.intersection(MATURITY_DATE_HEADERS):
        missing.add('APPRMATDT hoặc DSBSMATDT')
    return sorted(missing)



def loan_report(from_date: date, to_date: date) -> dict[str, Any]:
    table = _ensure_loan_table()
    if table is None:
        empty = {'total': 0, 'individual': 0, 'legal_entity': 0, 'closed': 0, 'active': 0}
        movements = {'total': 0, 'individual': 0, 'legal_entity': 0}
        return {
            'from_date': from_date.isoformat(), 'to_date': to_date.isoformat(),
            'from': empty.copy(), 'to': empty.copy(), 'change': empty.copy(),
            'increase': movements.copy(), 'decrease': movements.copy(), 'branches': [],
        }

    date_key = _key_for_header('DSBSDT')
    maturity_key = next((_key_for_header(header) for header in MATURITY_DATE_HEADERS if _key_for_header(header)), None)
    customer_key = _key_for_header('CUSTSEQ')
    type_key = _key_for_header('CUSTOMER_TYPE_CODE')
    branch_key = _key_for_header('BRCD')
    amount_keys = [_key_for_header(header) for header in (
        'DU_NO', 'PASTDUE_INTEREST_AMOUNT', 'TOTAL_INTEREST_REPAY_AMOUNT',
    )]
    if not all(key and key in table.c for key in [date_key, customer_key, type_key, branch_key, *amount_keys]):
        empty = {'total': 0, 'individual': 0, 'legal_entity': 0, 'closed': 0, 'active': 0}
        return {
            'from_date': from_date.isoformat(), 'to_date': to_date.isoformat(),
            'from': empty.copy(), 'to': empty.copy(),
            'change': empty.copy(), 'increase': {'total': 0, 'individual': 0, 'legal_entity': 0},
            'decrease': {'total': 0, 'individual': 0, 'legal_entity': 0}, 'branches': [],
        }

    date_col = cast(table.c[date_key], Date)
    cust_col = table.c[customer_key]
    type_col = table.c[type_key]
    branch_col = func.coalesce(table.c[branch_key], 'Chưa xác định').label('branch')

    def amount_expression(key: str):
        raw_amount = func.regexp_replace(cast(table.c[key], String), '[^0-9.-]', '', 'g')
        return cast(func.nullif(raw_amount, ''), Numeric)

    is_closed_cond = and_(*(func.coalesce(amount_expression(key), 0) == 0 for key in amount_keys))
    maturity_date_col = cast(table.c[maturity_key], Date) if maturity_key and maturity_key in table.c else None

    def matured_loan_cond(cutoff: date):
        if maturity_date_col is None:
            return false()
        return and_(
            is_closed_cond,
            maturity_date_col.is_not(None),
            maturity_date_col <= cutoff,
        )

    def active_loan_cond(cutoff: date):
        active_by_maturity = or_(
            ~is_closed_cond,
            maturity_date_col.is_(None),
            maturity_date_col > cutoff,
        ) if maturity_date_col is not None else true()
        return and_(date_col <= cutoff, active_by_maturity)

    def build_metrics_sql(cutoff: date, prefix: str):
        time_cond = date_col <= cutoff
        active_cond = active_loan_cond(cutoff)
        closed_at_cutoff = and_(time_cond, matured_loan_cond(cutoff))
        
        return {
            f'{prefix}_total': func.count(distinct(case((active_cond, cust_col), else_=None))).label(f'{prefix}_total'),
            f'{prefix}_individual': func.count(distinct(case((and_(active_cond, type_col == '100'), cust_col), else_=None))).label(f'{prefix}_individual'),
            f'{prefix}_legal_entity': func.count(distinct(case((and_(active_cond, type_col.in_(['520', '530'])), cust_col), else_=None))).label(f'{prefix}_legal_entity'),
            f'{prefix}_closed': func.count(case((closed_at_cutoff, table.c.id), else_=None)).label(f'{prefix}_closed'),
            f'{prefix}_active': func.count(case((active_cond, table.c.id), else_=None)).label(f'{prefix}_active'),
        }

    from_metrics = build_metrics_sql(from_date, 'from')
    to_metrics = build_metrics_sql(to_date, 'to')

    query = select(
        branch_col,
        *from_metrics.values(),
        *to_metrics.values()
    ).group_by(branch_col)

    def build_transition_query(include_branch: bool):
        grouped_branch = branch_col.label('branch')
        branch_group = [grouped_branch] if include_branch else []
        customer_group = [*branch_group, cust_col, type_col]
        active_customers = select(
            *customer_group,
            func.max(case((active_loan_cond(from_date), 1), else_=0)).label('from_active'),
            func.max(case((active_loan_cond(to_date), 1), else_=0)).label('to_active'),
        ).where(date_col <= to_date).group_by(*customer_group).subquery()

        added = and_(active_customers.c.from_active == 0, active_customers.c.to_active == 1)
        removed = and_(active_customers.c.from_active == 1, active_customers.c.to_active == 0)
        metrics = {}
        for label, condition in (('increase', added), ('decrease', removed)):
            metrics[f'{label}_total'] = func.count(distinct(case((condition, active_customers.c[customer_key]), else_=None))).label(f'{label}_total')
            metrics[f'{label}_individual'] = func.count(distinct(case((and_(condition, active_customers.c[type_key] == '100'), active_customers.c[customer_key]), else_=None))).label(f'{label}_individual')
            metrics[f'{label}_legal_entity'] = func.count(distinct(case((and_(condition, active_customers.c[type_key].in_(['520', '530'])), active_customers.c[customer_key]), else_=None))).label(f'{label}_legal_entity')

        if include_branch:
            return select(active_customers.c.branch, *metrics.values()).group_by(active_customers.c.branch)
        return select(*metrics.values())

    with engine.connect() as connection:
        results = connection.execute(query).mappings().all()
        global_result = connection.execute(
            select(*from_metrics.values(), *to_metrics.values())
        ).mappings().one()
        branch_transitions = connection.execute(build_transition_query(include_branch=True)).mappings().all()
        global_transitions = connection.execute(build_transition_query(include_branch=False)).mappings().one()

    branch_rows = []
    global_from = {key: global_result[f'from_{key}'] or 0 for key in ('total', 'individual', 'legal_entity', 'closed', 'active')}
    global_to = {key: global_result[f'to_{key}'] or 0 for key in ('total', 'individual', 'legal_entity', 'closed', 'active')}
    transitions_by_branch = {row['branch']: row for row in branch_transitions}
    global_increase = {key: global_transitions[f'increase_{key}'] or 0 for key in ('total', 'individual', 'legal_entity')}
    global_decrease = {key: global_transitions[f'decrease_{key}'] or 0 for key in ('total', 'individual', 'legal_entity')}

    for row in results:
        b_name = row['branch']
        
        # Bóc tách số liệu của từng chi nhánh
        b_from = {k: row[f'from_{k}'] or 0 for k in global_from.keys()}
        b_to = {k: row[f'to_{k}'] or 0 for k in global_to.keys()}
        transition = transitions_by_branch.get(b_name, {})
        b_increase = {key: transition.get(f'increase_{key}', 0) or 0 for key in ('total', 'individual', 'legal_entity')}
        b_decrease = {key: transition.get(f'decrease_{key}', 0) or 0 for key in ('total', 'individual', 'legal_entity')}

        branch_rows.append({
            'branch_code': b_name,
            'branch_name': b_name,
            'from': b_from,
            'to': b_to,
            'change': {k: b_to[k] - b_from[k] for k in b_to.keys()},
            'increase': b_increase,
            'decrease': b_decrease,
        })

    return {
        'from_date': from_date.isoformat(),
        'to_date': to_date.isoformat(),
        'from': global_from,
        'to': global_to,
        'change': {k: global_to[k] - global_from[k] for k in global_to.keys()},
        'increase': global_increase,
        'decrease': global_decrease,
        'branches': branch_rows,
    }



def loan_report_csv(report: dict[str, Any]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['Mã CN', 'Chi nhánh', 'Cá nhân - Đầu kỳ', 'Cá nhân - Cuối kỳ', 'Cá nhân - Tăng', 'Cá nhân - Giảm', 'Pháp nhân - Đầu kỳ', 'Pháp nhân - Cuối kỳ', 'Pháp nhân - Tăng', 'Pháp nhân - Giảm', 'Tổng - Đầu kỳ', 'Tổng - Cuối kỳ', 'Tổng - Tăng', 'Tổng - Giảm'])
    for row in report['branches']:
        writer.writerow([
            row['branch_code'], row['branch_name'],
            row['from']['individual'], row['to']['individual'], row['increase']['individual'], row['decrease']['individual'],
            row['from']['legal_entity'], row['to']['legal_entity'], row['increase']['legal_entity'], row['decrease']['legal_entity'],
            row['from']['total'], row['to']['total'], row['increase']['total'], row['decrease']['total'],
        ])
    return output.getvalue()


def loan_report_xlsx(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Báo cáo tổng quát'

    navy = '243C57'
    wine = 'A81948'
    header_fill = PatternFill('solid', fgColor='EAF0F6')
    total_fill = PatternFill('solid', fgColor='FFF0F3')
    white_font = Font(color='FFFFFF', bold=True, size=14)
    header_font = Font(color='46576C', bold=True, size=10)
    total_font = Font(color='96193F', bold=True)
    thin_border = Border(*(Side(style='thin', color='D6DEE8') for _ in range(4)))
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    sheet.merge_cells('A1:K1')
    sheet['A1'] = 'BÁO CÁO BIẾN ĐỘNG KHÁCH HÀNG VAY VỐN'
    sheet['A1'].fill = PatternFill('solid', fgColor=navy)
    sheet['A1'].font = white_font
    sheet['A1'].alignment = center
    sheet.merge_cells('A2:K2')
    sheet['A2'] = f"Mốc đầu kỳ: {report['from_date']}    |    Mốc cuối kỳ: {report['to_date']}    |    Đơn vị tính: Khách hàng"
    sheet['A2'].alignment = center

    sheet.merge_cells('A4:A5')
    sheet.merge_cells('B4:B5')
    sheet.merge_cells('C4:E4')
    sheet.merge_cells('F4:H4')
    sheet.merge_cells('I4:K4')
    headers = {'A4': 'Mã CN', 'B4': 'Chi nhánh', 'C4': 'CÁ NHÂN', 'F4': 'PHÁP NHÂN', 'I4': 'TỔNG'}
    for cell, value in headers.items():
        sheet[cell] = value
    subheaders = [
        f"Đầu kỳ ({report['from_date']})", f"Cuối kỳ ({report['to_date']})", 'Tăng / giảm',
        f"Đầu kỳ ({report['from_date']})", f"Cuối kỳ ({report['to_date']})", 'Tăng / giảm',
        f"Đầu kỳ ({report['from_date']})", f"Cuối kỳ ({report['to_date']})", 'Tăng / giảm',
    ]
    for index, value in enumerate(subheaders, start=3):
        sheet.cell(row=5, column=index, value=value)
    for row in sheet.iter_rows(min_row=4, max_row=5, min_col=1, max_col=11):
        for cell in row:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = thin_border

    for index, branch in enumerate(report.get('branches', []), start=1):
        row = [
            branch['branch_code'], branch['branch_name'],
            branch['from']['individual'], branch['to']['individual'],
            f"+{branch['increase']['individual']:,} / -{branch['decrease']['individual']:,}",
            branch['from']['legal_entity'], branch['to']['legal_entity'],
            f"+{branch['increase']['legal_entity']:,} / -{branch['decrease']['legal_entity']:,}",
            branch['from']['total'], branch['to']['total'],
            f"+{branch['increase']['total']:,} / -{branch['decrease']['total']:,}",
        ]
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=5 + index, column=column, value=value)
            cell.border = thin_border
            cell.alignment = center if column != 2 else Alignment(horizontal='left', vertical='center')
            if column >= 3:
                cell.number_format = '#,##0;[Red]-#,##0'
            if column in (5, 8, 11):
                cell.font = Font(color='46576C', bold=True)

    total_row = 6 + len(report.get('branches', []))
    total_values = [
        'TOÀN ĐƠN VỊ (đếm không trùng)', '',
        report['from']['individual'], report['to']['individual'],
        f"+{report['increase']['individual']:,} / -{report['decrease']['individual']:,}",
        report['from']['legal_entity'], report['to']['legal_entity'],
        f"+{report['increase']['legal_entity']:,} / -{report['decrease']['legal_entity']:,}",
        report['from']['total'], report['to']['total'],
        f"+{report['increase']['total']:,} / -{report['decrease']['total']:,}",
    ]
    for column, value in enumerate(total_values, start=1):
        cell = sheet.cell(row=total_row, column=column, value=value)
        cell.fill = total_fill
        cell.font = total_font
        cell.border = thin_border
        cell.alignment = center if column != 1 else Alignment(horizontal='left', vertical='center')
        if column in (3, 4, 6, 7, 9, 10):
            cell.number_format = '#,##0;[Red]-#,##0'

    widths = {'A': 8, 'B': 28, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 15, 'I': 15, 'J': 15, 'K': 15}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[4].height = 25
    sheet.row_dimensions[5].height = 34
    sheet.freeze_panes = 'C6'
    sheet.auto_filter.ref = f'A5:K{total_row}'

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _row_for_api(row: dict[str, Any], activity: dict[str, Any] | None, days_notice: int | None) -> dict[str, Any]:
    du_no_key = _key_for_header('DU_NO')
    resolved_status = (activity or {}).get('status') or row.get('status') or determine_status(row.get(du_no_key))
    result = dict(row)
    result['status'] = resolved_status
    result['gdv'] = (activity or {}).get('gdv_name') or (activity or {}).get('gdv') or ''
    result['gdv_note'] = (activity or {}).get('gdv_note') if (activity or {}).get('gdv_note') is not None else row.get('gdv_note') or ''
    due_value = row.get(loan_due_date_column) if loan_due_date_column else None
    parsed_due = _parse_date(due_value)
    result['days_to_due'] = (parsed_due - date.today()).days if parsed_due else None
    for header, key in zip(original_headers, column_keys):
        result[key] = _format_value(header, result.get(key))
    return result


def list_loans(page: int = 1, page_size: int = 50, days_notice: int | None = None, customer_code: str = '', customer_name: str = '', status: str = '', loan_category: str = '', username: str = '', role: str = '') -> dict[str, Any]:
    table = _ensure_loan_table()
    if table is None:
        return {'columns': [], 'loans': [], 'total': 0, 'page': page, 'page_size': page_size, 'total_pages': 0}

    code_key = _key_for_header('CUSTSEQ')
    name_key = _key_for_header('CUSTNM')
    with engine.connect() as connection:
        base_query = select(table)
        status_expression = _resolved_status_expression(table)
        if activity_table is not None and loan_dsbsseq_column is not None:
            base_query = base_query.select_from(
                table.outerjoin(activity_table, table.c[loan_dsbsseq_column] == activity_table.c.dsbsseq)
            )

        if customer_code and code_key and code_key in table.c:
            normalized_match = _normalized_like_expression(table.c[code_key], customer_code)
            if normalized_match is not None:
                base_query = base_query.where(normalized_match)
        if customer_name and name_key and name_key in table.c:
            normalized_match = _normalized_like_expression(table.c[name_key], customer_name)
            if normalized_match is not None:
                base_query = base_query.where(normalized_match)
        if status:
            base_query = base_query.where(status_expression == status)
        if days_notice is not None:
            due_column = table.c[loan_due_date_column] if loan_due_date_column and loan_due_date_column in table.c else None
            if due_column is None:
                base_query = base_query.where(false())
            else:
                due_date = cast(due_column, Date)
                base_query = base_query.where(
                    and_(
                        due_date >= func.current_date(),
                        due_date <= func.current_date() + days_notice,
                    )
                )
        if loan_category in {'overdue', 'due'}:
            if loan_category == 'overdue':
                due_date = _maturity_date_expression(table)
            else:
                due_column = table.c[loan_due_date_column] if loan_due_date_column and loan_due_date_column in table.c else None
                due_date = cast(due_column, Date) if due_column is not None else None
            if due_date is None:
                base_query = base_query.where(false())
            else:
                if loan_category == 'overdue':
                    base_query = base_query.where(
                        and_(
                            due_date < func.current_date(),
                            status_expression != 'Đã tất toán',
                        )
                    )
                else:
                    base_query = base_query.where(
                        and_(
                            due_date >= func.current_date(),
                            due_date <= func.current_date() + 5,
                        )
                    )
        elif loan_category in {'renewed', 'closed'}:
            category_status = 'Đã gia hạn' if loan_category == 'renewed' else 'Đã tất toán'
            base_query = base_query.where(status_expression == category_status)
        elif loan_category == 'edited':
            if activity_table is None or loan_dsbsseq_column is None:
                base_query = base_query.where(false())
            else:
                base_query = base_query.where(activity_table.c.gdv.is_not(None))
                if role.upper() == 'GDV':
                    base_query = base_query.where(activity_table.c.gdv == username)

        count_query = select(func.count()).select_from(base_query.subquery())
        total = connection.execute(count_query).scalar() or 0

        if total == 0:
            return _build_empty_response(page, page_size, _columns_for_api(days_notice))

        sort_column = table.c.id
        if loan_category == 'due' and loan_due_date_column and loan_due_date_column in table.c:
            sort_column = cast(table.c[loan_due_date_column], Date).asc().nullslast()
        paginated_query = base_query.order_by(sort_column, table.c.id).offset((page - 1) * page_size).limit(page_size)
        rows = connection.execute(paginated_query).mappings().fetchmany(page_size)

        activity_keys = {
            str(row[loan_dsbsseq_column]).strip()
            for row in rows
            if loan_dsbsseq_column and row.get(loan_dsbsseq_column) is not None and str(row[loan_dsbsseq_column]).strip()
        }
        activities = _activities_by_keys(connection, activity_keys)

        page_rows = []
        summary = _calculate_global_summary(connection, table)
        
        for raw in rows:
            row = dict(raw)
            activity_key = str(row[loan_dsbsseq_column]).strip() if loan_dsbsseq_column and row.get(loan_dsbsseq_column) is not None else ''
            activity = activities.get(activity_key)
            
            item = _row_for_api(row, activity, days_notice)
            page_rows.append(item)

        columns = _columns_for_api(days_notice)
        return {
            'columns': [{'label': label, 'key': key} for label, key in columns],
            'loans': page_rows,
            'total': total,
            'page': page,
            'page_size': page_size,
            'total_pages': (total + page_size - 1) // page_size,
            'summary': summary,
        }


def _columns_for_api(days_notice: int | None) -> list[tuple[str, str]]:
    columns = [('STT', '__serial__')]
    seen_labels: set[str] = set()
    for header, key in zip(original_headers, column_keys):
        if header.upper() not in MAPPED_HEADER_KEYS:
            continue
        label = HEADER_LABELS.get(header.upper(), header)
        if label not in seen_labels:
            columns.append((label, key))
            seen_labels.add(label)
        if days_notice is not None and key == loan_due_date_column:
            columns.append(('Số ngày còn lại', 'days_to_due'))

    columns.extend([('Trạng thái', 'status'), ('GDV', 'gdv')])
    if days_notice is not None and loan_due_date_column not in column_keys:
        columns.append(('Số ngày còn lại', 'days_to_due'))
    columns.append(('Ghi chú', 'gdv_note'))
    return columns


def _build_empty_response(page, page_size, columns=None):
    return {'columns': [{'label': label, 'key': key} for label, key in (columns or [])], 'loans': [], 'total': 0, 'page': page, 'page_size': page_size, 'total_pages': 0, 'summary': _empty_summary()}


def _empty_summary() -> dict[str, int | float]:
    return {
        'total': 0, 'due_5_days': 0, 'overdue': 0, 'renewed': 0, 'closed': 0,
        'total_outstanding': 0, 'individual_outstanding': 0, 'legal_entity_outstanding': 0,
    }

def _calculate_global_summary(connection, table):
    status_expression = _resolved_status_expression(table)
    outstanding = _outstanding_amount_expression(table)
    customer_type = _customer_type_expression(table)
    due_column = table.c[loan_due_date_column] if loan_due_date_column and loan_due_date_column in table.c else None
    due_date = cast(due_column, Date) if due_column is not None else None
    overdue_date = _maturity_date_expression(table)
    summary_query = select(
        func.count().label('total'),
        func.sum(case((status_expression == 'Đã tất toán', 1), else_=0)).label('closed'),
        func.sum(case((status_expression == 'Đã gia hạn', 1), else_=0)).label('renewed'),
        func.coalesce(func.sum(outstanding), 0).label('total_outstanding'),
    )
    if customer_type is not None:
        summary_query = summary_query.add_columns(
            func.coalesce(func.sum(case((customer_type == 'individual', outstanding), else_=0)), 0).label('individual_outstanding'),
            func.coalesce(func.sum(case((customer_type == 'legal_entity', outstanding), else_=0)), 0).label('legal_entity_outstanding'),
        )
    else:
        summary_query = summary_query.add_columns(
            func.cast(0, Numeric).label('individual_outstanding'),
            func.cast(0, Numeric).label('legal_entity_outstanding'),
        )
    if activity_table is not None and loan_dsbsseq_column is not None:
        summary_query = summary_query.select_from(
            table.outerjoin(activity_table, table.c[loan_dsbsseq_column] == activity_table.c.dsbsseq)
        )
    if due_date is not None:
        summary_query = summary_query.add_columns(
            func.sum(case((and_(due_date >= func.current_date(), due_date <= func.current_date() + 5), 1), else_=0)).label('due_5_days'),
            func.sum(case((and_(overdue_date < func.current_date(), status_expression != 'Đã tất toán'), 1), else_=0)).label('overdue'),
        )
    else:
        summary_query = summary_query.add_columns(func.cast(0, Integer).label('due_5_days'), func.cast(0, Integer).label('overdue'))
    result = connection.execute(summary_query).mappings().one()
    summary = _empty_summary()
    for key in summary:
        summary[key] = float(result.get(key) or 0) if key.endswith('outstanding') else int(result.get(key) or 0)
    return summary

def get_loan(loan_id: int) -> dict[str, Any]:
    table = _ensure_loan_table()
    if table is None:
        raise HTTPException(status_code=404, detail='Chưa có dữ liệu khoản vay.')
    with engine.connect() as connection:
        row = connection.execute(select(table).where(table.c.id == loan_id)).mappings().first()
        if not row:
            raise HTTPException(status_code=404, detail='Không tìm thấy khoản vay.')
        activity = _activity_by_key(connection, str(row.get(loan_dsbsseq_column)).strip() if loan_dsbsseq_column and row.get(loan_dsbsseq_column) is not None else None)
        return _row_for_api(dict(row), activity, None)


def get_loan_detail_data(loan_id: int) -> dict[str, Any]:
    loan = get_loan(loan_id)
    columns = [
        {'label': HEADER_LABELS.get(header.upper(), header), 'key': key}
        for header, key in zip(original_headers, column_keys)
    ]
    return {'loan': loan, 'columns': columns}


def update_loan(loan_id: int, status: str, note: str, gdv_username: str) -> dict[str, Any]:
    table = _ensure_loan_table()
    activity = _ensure_activity_table()
    if table is None:
        raise HTTPException(status_code=404, detail='Chưa có dữ liệu khoản vay.')
    with engine.begin() as connection:
        row = connection.execute(select(table).where(table.c.id == loan_id)).mappings().first()
        if not row:
            raise HTTPException(status_code=404, detail='Không tìm thấy khoản vay.')
        connection.execute(update(table).where(table.c.id == loan_id).values(status=status, gdv_note=note.strip()))
        key = str(row.get(loan_dsbsseq_column)).strip() if loan_dsbsseq_column and row.get(loan_dsbsseq_column) is not None else ''
        if key:
            existing = connection.execute(select(activity).where(activity.c.dsbsseq == key)).mappings().first()
            payload = {'dsbsseq': key, 'status': status, 'gdv_note': note.strip(), 'gdv': gdv_username, 'updated_at': date.today().isoformat()}
            if existing:
                connection.execute(update(activity).where(activity.c.id == existing['id']).values(**payload))
            else:
                connection.execute(insert(activity), payload)
    return get_loan(loan_id)


def _read_upload(file: UploadFile) -> pd.DataFrame:
    filename = (file.filename or '').lower()
    if not filename.endswith(('.csv', '.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail='Chỉ hỗ trợ file CSV, XLSX hoặc XLS.')
    try:
        content = file.file.read()
        return pd.read_csv(io.BytesIO(content), encoding='utf-8-sig') if filename.endswith('.csv') else pd.read_excel(io.BytesIO(content))
    except Exception as error:
        raise HTTPException(status_code=400, detail=f'Không thể đọc file {file.filename}: {error}') from error


def import_loan_files(files: list[UploadFile]) -> dict[str, Any]:
    frames: list[tuple[str, pd.DataFrame]] = []
    headers: list[str] | None = None
    for file in files:
        frame = _read_upload(file)
        if frame.empty:
            continue
        current = [normalize_header(column) for column in frame.columns]
        missing = _missing_import_columns(current)
        if missing:
            raise HTTPException(status_code=400, detail=f'File {file.filename} thiếu cột bắt buộc: {", ".join(missing)}')
        if headers is None:
            headers = current
        elif current != headers:
            raise HTTPException(status_code=400, detail=f'File {file.filename} có cấu trúc khác với file trước đó.')
        frames.append((file.filename or 'unknown', frame))
    if not frames or headers is None:
        raise HTTPException(status_code=400, detail='Không có dữ liệu hợp lệ trong các file đã chọn.')
    table = _ensure_loan_table(headers)
    _ensure_activity_table()
    dsbsseq_header = next((header for header in headers if header.upper() == 'DSBSSEQ'), None)
    du_no_header = next((header for header in headers if header.upper() == 'DU_NO'), None)
    imported = 0
    with engine.begin() as connection:
        connection.execute(delete(table))
        for filename, frame in frames:
            batch = []
            for source_values in frame.itertuples(index=False, name=None):
                source_row = dict(zip(headers, source_values))
                data = {key: (None if pd.isna(source_row.get(header)) else str(source_row.get(header))) for header, key in zip(headers, column_keys)}
                du_no = source_row.get(du_no_header) if du_no_header else None
                data.update({'status': determine_status(du_no), 'gdv_note': '', 'source_file': os.path.basename(filename), 'imported_at': date.today().isoformat()})
                batch.append(data)
                if len(batch) >= BATCH_SIZE:
                    _restore_activity_state(connection, batch)
                    connection.execute(insert(table), batch)
                    imported += len(batch)
                    batch.clear()
            if batch:
                _restore_activity_state(connection, batch)
                connection.execute(insert(table), batch)
                imported += len(batch)
    return {'message': f'Đã import thành công {len(frames)} file.', 'rows_imported': imported, 'files_imported': len(frames)}


def _restore_activity_state(connection, batch: list[dict[str, Any]]) -> None:
    if not loan_dsbsseq_column or activity_table is None:
        return
    keys = {str(row[loan_dsbsseq_column]).strip() for row in batch if row.get(loan_dsbsseq_column)}
    activities = _activities_by_keys(connection, keys)
    for row in batch:
        activity = activities.get(str(row[loan_dsbsseq_column]).strip()) if row.get(loan_dsbsseq_column) else None
        if activity:
            row['status'] = activity.get('status') or row['status']
            row['gdv_note'] = activity.get('gdv_note') or ''
